"""
使用python-pptx生成路演PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_ppt():
    """创建路演PPT"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 幻灯片1：封面
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(12.333), Inches(1.5))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "基于上下文感知与自洽性验证的\n长文档智能标注系统"
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER

    # 副标题
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = "FlagOS开放计算全球挑战赛 - 赛道三"
    subtitle_para.font.size = Pt(28)
    subtitle_para.alignment = PP_ALIGN.CENTER

    # 日期
    date_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.333), Inches(0.5))
    date_frame = date_box.text_frame
    date_para = date_frame.paragraphs[0]
    date_para.text = "2026年2月"
    date_para.font.size = Pt(18)
    date_para.alignment = PP_ALIGN.CENTER

    # 幻灯片2：项目背景
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "项目背景")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    points = [
        "FlagOS开放计算全球挑战赛 - 赛道三：自动数据标注",
        "核心挑战：长上下文技术文档的结构化信息抽取",
        "技术约束：必须使用Qwen3-4B模型，禁止微调",
        "解决方案：基于In-Context Learning (ICL)范式",
        "目标：实现高质量、自动化的文档标注系统"
    ]

    for i, point in enumerate(points):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = "• " + point
        p.font.size = Pt(24)
        p.space_before = Pt(20)

    # 幻灯片3：技术架构
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "技术架构")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    modules = [
        "数据预处理模块 → 文档清洗、分块、示例库构建",
        "上下文构建模块 → 动态构建摘要、章节路径、前后文",
        "示例检索模块 → 基于语义相似度检索Few-shot示例",
        "提示词工程模块 → 可配置的提示词模板管理",
        "自洽性验证模块 → 多轮验证确保标注质量",
        "后处理模块 → 结构化输出解析、清洗、去重"
    ]

    for i, module in enumerate(modules):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = f"{i+1}. {module}"
        p.font.size = Pt(22)
        p.space_before = Pt(15)

    # 幻灯片4：核心创新
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "核心创新点")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    innovations = [
        "分层级上下文感知机制",
        "  - 文档摘要 + 章节路径 + 前后文动态组装",
        "  - 有效处理长文档，控制token消耗",
        "动态Few-shot学习",
        "  - 基于语义相似度检索最相关示例",
        "  - 提高标注一致性和准确性",
        "自洽性验证机制",
        "  - 多轮验证确保结果完整准确",
        "  - 自动修正遗漏和错误"
    ]

    for i, item in enumerate(innovations):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(20)
        p.space_before = Pt(10)

    # 幻灯片5：技术栈
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "技术栈")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    tech_stack = [
        "核心模型：Qwen3-4B (FlagOS平台)",
        "嵌入模型：BAAI/bge-m3",
        "编程语言：Python 3.8+",
        "主要框架：PyTorch, Transformers",
        "日志系统：Loguru",
        "配置管理：YAML"
    ]

    for i, tech in enumerate(tech_stack):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = "✓ " + tech
        p.font.size = Pt(24)
        p.space_before = Pt(18)

    # 幻灯片6：评测指标
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "评测指标")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    metrics = [
        "精确率 (Precision) = 正确预测数 / 总预测数",
        "召回率 (Recall) = 正确预测数 / 真实标签数",
        "F1分数 = 2 × P × R / (P + R)",
        "实体识别准确率",
        "关系抽取准确率"
    ]

    for i, metric in enumerate(metrics):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = "• " + metric
        p.font.size = Pt(24)
        p.space_before = Pt(20)

    # 幻灯片7：项目亮点
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "项目亮点")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    highlights = [
        "完全基于ICL范式，无需模型微调",
        "模块化设计，各组件独立可配置",
        "完整的代码和配置，确保结果可复现",
        "支持多种实体类型和关系类型",
        "自洽性验证确保标注质量"
    ]

    for i, highlight in enumerate(highlights):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = "★ " + highlight
        p.font.size = Pt(24)
        p.space_before = Pt(20)

    # 幻灯片8：总结
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "总结与展望")

    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    summary = [
        "本项目实现了基于ICL的长文档智能标注系统",
        "通过分层上下文感知有效处理长文档",
        "动态Few-shot提高标注一致性",
        "自洽性验证确保标注质量",
        "",
        "后续优化方向：",
        "  - 优化上下文长度控制策略",
        "  - 扩展更多实体和关系类型"
    ]

    for i, item in enumerate(summary):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(22)
        p.space_before = Pt(12)

    # 幻灯片9：致谢
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "感谢聆听"
    title_para.font.size = Pt(48)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = "FlagOS开放计算全球挑战赛"
    subtitle_para.font.size = Pt(24)
    subtitle_para.alignment = PP_ALIGN.CENTER

    prs.save('c:/D/compet/dcic/FlagOS开放计算全球挑战赛/output/ppt/路演PPT.pptx')
    print('路演PPT生成完成！')


def add_title(slide, title_text):
    """添加幻灯片标题"""
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title_text
    title_para.font.size = Pt(36)
    title_para.font.bold = True


if __name__ == '__main__':
    create_ppt()
