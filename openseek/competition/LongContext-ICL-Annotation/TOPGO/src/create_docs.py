"""
使用python-docx生成Word文档
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_tech_doc():
    """创建技术方案文档"""
    doc = Document()
    
    # 标题
    title = doc.add_heading('基于上下文感知与自洽性验证的长文档智能标注系统', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_heading('技术方案文档', 1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph('FlagOS开放计算全球挑战赛 - 赛道三')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph('文档版本：1.0 | 最后更新：2026-02-19')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # 第一章
    doc.add_heading('一、项目概述', 1)
    doc.add_heading('1.1 项目背景', 2)
    doc.add_paragraph('本项目为FlagOS开放计算全球挑战赛赛季一的"自动数据标注"赛道参赛方案。赛道要求参赛者基于Qwen3-4B模型，使用组委会提供的统一数据集，设计有效的In-Context Learning (ICL)方案进行自动数据标注。')
    
    doc.add_heading('1.2 核心目标', 2)
    doc.add_paragraph('开发一套自动化系统，能够对长上下文技术文档（如学术论文、技术报告）进行高质量、结构化的信息抽取与标注。系统需完全基于Qwen3-4B模型，采用ICL范式，不进行模型微调。')
    
    doc.add_heading('1.3 设计原则', 2)
    doc.add_paragraph('合规性：严格使用指定的Qwen3-4B模型，不使用任何违规的模型增强或第三方服务。')
    doc.add_paragraph('高效性：针对长文档特性进行优化，控制计算与token消耗。')
    doc.add_paragraph('鲁棒性：能够处理多样化的文档结构和内容。')
    doc.add_paragraph('可复现性：提供完整的代码和配置，确保结果可复现。')
    
    # 第二章
    doc.add_heading('二、技术架构', 1)
    doc.add_heading('2.1 系统架构图', 2)
    doc.add_paragraph('系统采用模块化设计，各组件独立可配置：')
    doc.add_paragraph('[原始数据集] → [数据预处理模块] → [示例库]')
    doc.add_paragraph('[待标注文档] → [分层上下文构建模块] → [优化上下文]')
    doc.add_paragraph('[优化上下文] + [动态示例检索] → [提示词组装模块]')
    doc.add_paragraph('[提示词组装] → [Qwen3-4B模型] → [自洽性验证模块]')
    doc.add_paragraph('[自洽性验证] → [后处理与输出模块] → [最终标注结果]')
    
    doc.add_heading('2.2 核心模块说明', 2)
    doc.add_paragraph('1. 数据预处理模块：负责文档清洗、分块、示例库构建。')
    doc.add_paragraph('2. 上下文构建模块：动态构建文档摘要、章节路径、前后文上下文。')
    doc.add_paragraph('3. 示例检索模块：基于语义相似度检索最相关的Few-shot示例。')
    doc.add_paragraph('4. 提示词工程模块：可配置的提示词模板和管理器。')
    doc.add_paragraph('5. 自洽性验证模块：多轮验证确保标注结果的准确性和完整性。')
    doc.add_paragraph('6. 后处理模块：结构化输出解析、清洗、去重。')
    doc.add_paragraph('7. 评测模块：计算Precision、Recall、F1等指标。')
    
    # 第三章
    doc.add_heading('三、核心算法', 1)
    doc.add_heading('3.1 分层级上下文构建算法', 2)
    doc.add_paragraph('算法流程：')
    doc.add_paragraph('1. 文档摘要生成：调用Qwen3-4B，生成约150-200字的文档摘要。')
    doc.add_paragraph('2. 章节结构解析：基于Markdown格式识别章节标题，建立文档大纲。')
    doc.add_paragraph('3. 目标定位与上下文裁剪：根据标注任务确定相关段落，组装上下文。')
    
    doc.add_heading('3.2 动态Few-shot检索算法', 2)
    doc.add_paragraph('使用BGE-M3嵌入模型对示例库和待标注文本进行向量化，通过余弦相似度计算选取Top-K个最相关的示例。')
    
    doc.add_heading('3.3 自洽性验证算法', 2)
    doc.add_paragraph('设计"质检员"角色的验证提示词，将初始标注结果、原始任务指令和原始上下文片段再次提交给Qwen3-4B进行验证。')
    
    # 第四章
    doc.add_heading('四、实现方案', 1)
    doc.add_heading('4.1 技术栈', 2)
    doc.add_paragraph('编程语言：Python 3.8+')
    doc.add_paragraph('核心模型：Qwen3-4B')
    doc.add_paragraph('嵌入模型：BAAI/bge-m3')
    doc.add_paragraph('主要依赖：torch, transformers, sentence-transformers, loguru')
    
    doc.add_heading('4.2 项目结构', 2)
    doc.add_paragraph('src/ - 源代码目录')
    doc.add_paragraph('  data/ - 数据处理模块')
    doc.add_paragraph('  models/ - 模型调用模块')
    doc.add_paragraph('  utils/ - 工具模块')
    doc.add_paragraph('  context_builder.py - 上下文构建器')
    doc.add_paragraph('  retriever.py - 示例检索器')
    doc.add_paragraph('  prompt_engineer.py - 提示词工程')
    doc.add_paragraph('  validator.py - 自洽性验证器')
    doc.add_paragraph('  postprocessor.py - 后处理器')
    doc.add_paragraph('  evaluator.py - 评测器')
    doc.add_paragraph('  pipeline.py - 主Pipeline')
    
    # 第五章
    doc.add_heading('五、测试方案', 1)
    doc.add_heading('5.1 功能测试', 2)
    doc.add_paragraph('数据加载测试：验证JSONL/JSON格式数据正确解析。')
    doc.add_paragraph('文档分块测试：验证长文档正确分块。')
    doc.add_paragraph('示例检索测试：验证相似度检索准确性。')
    doc.add_paragraph('标注输出测试：验证JSON格式输出正确。')
    
    doc.add_heading('5.2 性能测试', 2)
    doc.add_paragraph('单文档处理时间：确保在算力限制内完成。')
    doc.add_paragraph('内存占用测试：验证内存使用可控。')
    doc.add_paragraph('并发处理测试：验证批量处理能力。')
    
    doc.add_heading('5.3 评测指标', 2)
    doc.add_paragraph('精确率(Precision)：正确预测数/总预测数')
    doc.add_paragraph('召回率(Recall)：正确预测数/真实标签数')
    doc.add_paragraph('F1分数：精确率和召回率的调和平均')
    
    # 第六章
    doc.add_heading('六、总结与展望', 1)
    doc.add_heading('6.1 项目亮点', 2)
    doc.add_paragraph('创新的分层级上下文感知机制，有效处理长文档。')
    doc.add_paragraph('动态Few-shot学习，提高标注一致性。')
    doc.add_paragraph('自洽性验证机制，确保标注质量。')
    doc.add_paragraph('完全模块化设计，便于调试和优化。')
    
    doc.add_heading('6.2 后续优化方向', 2)
    doc.add_paragraph('优化上下文长度控制策略，提高长文档处理能力。')
    doc.add_paragraph('扩展更多实体类型和关系类型。')
    doc.add_paragraph('探索更高效的示例检索算法。')
    
    doc.save('c:/D/compet/dcic/FlagOS开放计算全球挑战赛/output/docs/技术方案文档.docx')
    print('技术方案文档生成完成！')


def create_manual_doc():
    """创建使用说明书"""
    doc = Document()
    
    # 标题
    title = doc.add_heading('基于上下文感知与自洽性验证的长文档智能标注系统', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_heading('使用说明书', 1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph('文档版本：1.0 | 最后更新：2026-02-19')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # 第一章
    doc.add_heading('一、系统概述', 1)
    doc.add_paragraph('本系统是一个基于In-Context Learning (ICL)范式的长文档智能标注系统，用于从技术文档中提取结构化信息。系统完全基于Qwen3-4B模型，无需模型微调即可实现高质量标注。')
    
    # 第二章
    doc.add_heading('二、环境准备', 1)
    doc.add_heading('2.1 系统要求', 2)
    doc.add_paragraph('操作系统：Windows/Linux/MacOS')
    doc.add_paragraph('Python版本：3.8+')
    doc.add_paragraph('内存：建议8GB以上')
    doc.add_paragraph('存储：至少2GB可用空间')
    
    doc.add_heading('2.2 安装步骤', 2)
    doc.add_paragraph('步骤1：创建虚拟环境')
    doc.add_paragraph('  python -m venv venv')
    doc.add_paragraph('  source venv/bin/activate  # Linux/Mac')
    doc.add_paragraph('  venv\\Scripts\\activate    # Windows')
    doc.add_paragraph('步骤2：安装依赖')
    doc.add_paragraph('  pip install -r requirements.txt')
    doc.add_paragraph('步骤3：配置API')
    doc.add_paragraph('  编辑configs/config.yaml，设置模型API端点')
    
    # 第三章
    doc.add_heading('三、配置说明', 1)
    doc.add_heading('3.1 模型配置', 2)
    doc.add_paragraph('model:')
    doc.add_paragraph('  name: Qwen3-4B        # 模型名称')
    doc.add_paragraph('  api_base: API端点     # 比赛平台提供')
    doc.add_paragraph('  api_key: 您的密钥      # API访问密钥')
    doc.add_paragraph('  temperature: 0.1        # 生成温度')
    doc.add_paragraph('  max_tokens: 2048        # 最大输出长度')
    
    doc.add_heading('3.2 上下文配置', 2)
    doc.add_paragraph('context:')
    doc.add_paragraph('  summary_length: 200     # 摘要长度')
    doc.add_paragraph('  neighbor_paragraphs_before: 1  # 前N段')
    doc.add_paragraph('  neighbor_paragraphs_after: 1   # 后M段')
    doc.add_paragraph('  max_context_length: 8000       # 最大上下文长度')
    
    doc.add_heading('3.3 检索配置', 2)
    doc.add_paragraph('retrieval:')
    doc.add_paragraph('  embedding_model: BAAI/bge-m3  # 嵌入模型')
    doc.add_paragraph('  top_k: 3                          # Top-K示例数')
    doc.add_paragraph('  similarity_threshold: 0.5         # 相似度阈值')
    
    # 第四章
    doc.add_heading('四、数据格式', 1)
    doc.add_heading('4.1 输入数据格式', 2)
    doc.add_paragraph('训练数据 (data/raw/train.jsonl)：')
    doc.add_paragraph('{ doc_id: 文档ID, text: 文档内容, annotation: { entities: [...], relations: [...] } }')
    
    doc.add_heading('4.2 输出数据格式', 2)
    doc.add_paragraph('预测结果 (output/results/predictions.json)：')
    doc.add_paragraph('[{ doc_id: 文档ID, entities: [...], relations: [...] }]')
    
    # 第五章
    doc.add_heading('五、运行指南', 1)
    doc.add_heading('5.1 运行标注任务', 2)
    doc.add_paragraph('# 将数据放置到 data/raw/ 目录')
    doc.add_paragraph('# 运行标注')
    doc.add_paragraph('python scripts/run_annotation.py')
    
    doc.add_heading('5.2 运行评测', 2)
    doc.add_paragraph('python scripts/run_evaluation.py --predictions output/results/predictions.json --ground-truth data/raw/ground_truth.jsonl')
    
    doc.add_heading('5.3 运行演示', 2)
    doc.add_paragraph('# 使用模拟数据进行演示')
    doc.add_paragraph('python scripts/demo.py')
    
    # 第六章
    doc.add_heading('六、常见问题', 1)
    doc.add_paragraph('Q1: 如何配置API端点？')
    doc.add_paragraph('A: 编辑configs/config.yaml文件，设置model.api_base和model.api_key字段。')
    doc.add_paragraph('Q2: 如何处理超长文档？')
    doc.add_paragraph('A: 系统会自动进行文档分块，可通过调整context.chunk_size参数控制分块大小。')
    doc.add_paragraph('Q3: 如何添加新的实体类型？')
    doc.add_paragraph('A: 在configs/config.yaml的annotation.entity_types列表中添加新类型。')
    doc.add_paragraph('Q4: 标注结果格式错误怎么办？')
    doc.add_paragraph('A: 系统的后处理模块会自动尝试修复常见格式问题，如仍有问题请检查原始输出。')
    
    # 第七章
    doc.add_heading('七、技术支持', 1)
    doc.add_paragraph('如遇到问题，请参考以下资源：')
    doc.add_paragraph('比赛页面: https://modelscope.cn/competition/180')
    doc.add_paragraph('FlagOS官网: https://flagos.net/')
    doc.add_paragraph('GitHub: https://github.com/FlagOpen')
    
    doc.save('c:/D/compet/dcic/FlagOS开放计算全球挑战赛/output/docs/使用说明书.docx')
    print('使用说明书生成完成！')


if __name__ == '__main__':
    create_tech_doc()
    create_manual_doc()
